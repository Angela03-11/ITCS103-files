#Quejano_FinalProject

import tkinter as gela
import openpyxl as op
from tkinter import ttk, messagebox

window = gela.Tk()
window.title("Meeting Room Booking System")
window.geometry("1200x650")
window.resizable(False, False)
window.configure(bg="#DCEEFF", cursor = "hand2") #to do: #ffffff


def display():

    workbook = op.load_workbook("Quejano_Database.xlsx")
    sheet = workbook.active

    for item in table.get_children():
        table.delete(item)

    for row in sheet.iter_rows(min_row=2, values_only=True):
        table.insert("", gela.END, values=row)


def validation():
    name = name_entry.get()
    room = room_combo.get() # I seeked help po sa AI para sa feature na gusto ko which is ang magkaroon ng selection sa pag book ng room
    date = date_entry.get()
    start_time = time_entry.get()
    duration = duration_entry.get()
    
    if not name or not room or not date or not start_time or not duration:
        messagebox.showerror("Error", "Please fill in all fields.")
        return False
    
    if not duration.isdigit():
        messagebox.showerror("Error", "Duration must contain numbers only.")
        return False
    
    return True


def create():
    if not validation():
        return

    workbook = op.load_workbook("Quejano_Database.xlsx")
    sheet = workbook.active

    new_id = sheet.max_row 

    total_cost = compute_cost()

    sheet.append([new_id, name_entry.get(), room_combo.get(), date_entry.get(), time_entry.get(), duration_entry.get(), total_cost])

    workbook.save("Quejano_Database.xlsx")
    messagebox.showinfo("Success", "Booking Added Successfully!")

    display()


def auto_populate(event):

    selected = table.focus()

    values = table.item(selected, "values")

    if values:

        name_entry.delete(0, gela.END)
        date_entry.delete(0, gela.END)
        time_entry.delete(0, gela.END)
        duration_entry.delete(0, gela.END)

        name_entry.insert(0, values[1])
        room_combo.set(values[2])
        date_entry.insert(0, values[3])
        time_entry.insert(0, values[4])
        duration_entry.insert(0, values[5])

        cost_label.config(text = f"₱ {values[6]}")


def update():

    selected = table.focus()

    if not selected:
        messagebox.showerror("Error", "Select a record first.")
        return

    if not validation():
        return

    values = table.item(selected, "values")
    record_id = values[0]

    total_cost = compute_cost()

    workbook = op.load_workbook("Quejano_Database.xlsx")
    sheet = workbook.active

    for row in sheet.iter_rows(min_row=2):
        if str(record_id) == str(row[0].value):
            row[1].value = name_entry.get()
            row[2].value = room_combo.get()
            row[3].value = date_entry.get()
            row[4].value = time_entry.get()
            row[5].value = duration_entry.get()
            row[6].value = total_cost
            break

    workbook.save("Quejano_Database.xlsx")
    messagebox.showinfo("Success", "Booking Updated Successfully!")

    display()


def delete():

    selected = table.focus()

    if not selected:
        messagebox.showerror("Error", "Select a record first.")
        return

    confirm = messagebox.askyesnocancel("Confirm", "Are you sure you want to cancel this booking?")

    if not confirm:
        return

    values = table.item(selected, "values")
    record_id = values[0]

    workbook = op.load_workbook("Quejano_Database.xlsx")
    sheet = workbook.active

    for i, row in enumerate(
        sheet.iter_rows(min_row = 2), start = 2):

        if str(record_id) == str(row[0].value):
            sheet.delete_rows(i)
            break

    workbook.save("Quejano_Database.xlsx")
    messagebox.showinfo("Success", "Booking Deleted Successfully!")

    display()


def compute_cost():

    room = room_combo.get()
    duration = int(duration_entry.get())

    if room == "Room A":
        rate = 200

    elif room == "Room B":
        rate = 300

    elif room == "Room C":
        rate = 400

    elif room == "Room D":
        rate = 500
        
    elif room == "Room E":
        rate = 600
        
    elif room == "Executive Room":
        rate = 2000
        
    elif room == "Board Room":
        rate = 1500
        
    elif room == "Client Meeting Room":
        rate = 1000
        
    elif room == "Conference Room":
        rate = 1200
        
    elif room == "VIP Room":
        rate = 3000
    
    total = rate * duration

    return total


def show_cost(event):

    if duration_entry.get().isdigit():

        total = compute_cost()
        cost_label.config(text = f"₱ {total}")


#Title

title = gela.Label(window, text="MEETING ROOM BOOKING SYSTEM", font = ("Poppins", 20, "bold"), bg = "#DCEEFF", fg = "#003566")
title.place(x = 350, y = 15)


#Framee
form_frame = gela.Frame(window, bg = "white", bd = 2, relief = "groove")
form_frame.place(x = 20,y = 70, width = 300, height = 540)


#Name
label_name = gela.Label(form_frame, text = "Name", bg = "white", font = ("Poppins", 10, "bold"))
label_name.place(x = 20, y = 20)


#Name ent
name_entry = gela.Entry(form_frame, font = ("Poppins", 11), width = 28)
name_entry.place(x = 20, y = 45)


#Room
label_room = gela.Label(form_frame, text = "Meeting Room", bg = "white", font = ("Poppins", 10, "bold"))
label_room.place(x = 20, y = 85)


#Room ent 
room_combo = ttk.Combobox(form_frame, values = [ "Room A", "Room B", "Room C", "Room D",  "Room E", "Executive Room", "Board Room", "Client Meeting Room", "Conference Room", "VIP Room"], width = 25)
room_combo.place(x = 20, y = 110)  



#Date
label_date = gela.Label(form_frame, text = "Date", bg = "white", font = ("Poppins", 10, "bold"))
label_date.place(x = 20, y = 150)



#Date ent
date_entry = gela.Entry(form_frame, font = ("Poppins", 11), width = 28)
date_entry.place(x = 20, y = 175)



#Time
label_time = gela.Label(form_frame, text = "Start Time", bg = "white", font = ("Poppins", 10, "bold"))
label_time.place(x = 20, y = 215)



#Time ent
time_entry = gela.Entry(form_frame, font = ("Poppins", 11), width = 28)
time_entry.place(x = 20, y = 240)



#Time sample label
time_sample = gela.Label(form_frame, text = "Example: 9:30 AM", bg = "white", fg = "gray")
time_sample.place(x = 20, y = 265)



#Duration
label_duration = gela.Label(form_frame, text = "Duration (Hours)", bg = "white", font = ("Poppins", 10, "bold"))
label_duration.place(x = 20, y = 295)



#Duration ent
duration_entry = gela.Entry(form_frame, font = ("Poppins", 11), width = 28)
duration_entry.place(x = 20, y = 320)



#Bind
duration_entry.bind("<KeyRelease>", show_cost) # studied this part, para mag automatic show 'yung total cost ng booking


#Total cost label
label_cost = gela.Label(form_frame, text = "Total Cost", bg = "white", font = ("Poppins", 10, "bold"))
label_cost.place(x = 20, y = 360)



cost_label = gela.Label(form_frame, text = "₱ 0", bg = "white", fg = "green", font = ("Poppins", 16, "bold"))
cost_label.place(x = 20, y = 390)




#Add button
add_btn = gela.Button(form_frame, text = "ADD BOOKING", bg ="#90EE90", width = 22, command = create)   
add_btn.place(x = 55, y = 440)



#Update button
update_btn = gela.Button(form_frame, text = "UPDATE BOOKING", bg ="#FFD166", width = 22, command = update)
update_btn.place(x = 55, y = 475)



#Delete button
delete_btn = gela.Button(form_frame, text = "DELETE BOOKING", bg ="#FF6B6B", fg = "white", width = 22, command = delete)
delete_btn.place(x = 55, y = 510)



#Table with frame
table_frame = gela.Frame(window, bg = "white", bd = 2, relief = "groove")
table_frame.place(x = 340, y = 70, width = 840, height = 540)



#Treeview
table = ttk.Treeview(table_frame,columns = ("ID", "Name", "Room", "Date", "Start Time", "Duration", "Total Cost"), show = "headings", height = 22)

columns = ("ID", "Name", "Room", "Date", "Start Time", "Duration", "Total Cost")

for col in columns:

    table.heading(col, text = col)

table.column("ID", width = 50)
table.column("Name", width = 150)
table.column("Room", width = 150)
table.column("Date", width = 120)
table.column("Start Time", width = 120)
table.column("Duration", width = 100)
table.column("Total Cost", width = 120)

table.place(x= 10, y = 10, width = 820, height = 520)

table.bind("<<TreeviewSelect>>", auto_populate)

display()

window.mainloop()